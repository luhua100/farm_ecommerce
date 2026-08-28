from pathlib import Path
from shutil import copyfile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
APP_OUTPUT = ROOT / "assets" / "doc" / "privacy_policy.docx"
SERVICE_OUTPUT = (
    ROOT.parent / "productsfarm" / "Service" / "public" / "legal"
    / "privacy_policy.docx"
)

# Use one CJK-capable family for every run so LibreOffice and mobile viewers
# render Chinese consistently instead of falling back to missing glyph boxes.
FONT_ASCII = "Arial Unicode MS"
FONT_EAST_ASIA = "Arial Unicode MS"
INK = RGBColor(28, 37, 32)
MUTED = RGBColor(91, 105, 97)
HEADING = RGBColor(46, 116, 181)
LIGHT_FILL = "F2F4F7"


def set_run_font(run, size=11, color=INK, bold=False):
    run.font.name = FONT_ASCII
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_EAST_ASIA)
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.bold = bold


def add_page_number(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])
    set_run_font(run, size=9, color=MUTED)


def add_hyperlink(paragraph, text, url):
    relationship_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "087B4C")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_properties.extend([color, underline])
    run.append(run_properties)
    text_element = OxmlElement("w:t")
    text_element.text = text
    run.append(text_element)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def configure_styles(document):
    normal = document.styles["Normal"]
    normal.font.name = FONT_ASCII
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_EAST_ASIA)
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    heading_tokens = {
        "Heading 1": (16, 16, 8),
        "Heading 2": (13, 12, 6),
        "Heading 3": (12, 8, 4),
    }
    for style_name, (size, before, after) in heading_tokens.items():
        style = document.styles[style_name]
        style.font.name = FONT_ASCII
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_EAST_ASIA)
        style.font.size = Pt(size)
        style.font.color.rgb = HEADING
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Number"):
        style = document.styles[style_name]
        style.font.name = FONT_ASCII
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_EAST_ASIA)
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.167


def add_body(document, text, *, bold_prefix=None):
    paragraph = document.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        prefix = paragraph.add_run(bold_prefix)
        set_run_font(prefix, bold=True)
        body = paragraph.add_run(text[len(bold_prefix):])
        set_run_font(body)
    else:
        run = paragraph.add_run(text)
        set_run_font(run)
    return paragraph


def add_bullet(document, text):
    paragraph = document.add_paragraph(style="List Bullet")
    run = paragraph.add_run(text)
    set_run_font(run)
    return paragraph


def add_callout(document, title, text):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.18)
    paragraph.paragraph_format.right_indent = Inches(0.18)
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(10)
    paragraph.paragraph_format.line_spacing = 1.167
    properties = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), LIGHT_FILL)
    properties.append(shading)
    title_run = paragraph.add_run(f"{title}\n")
    set_run_font(title_run, bold=True)
    body_run = paragraph.add_run(text)
    set_run_font(body_run)


def build_document():
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    configure_styles(document)

    header = section.header.paragraphs[0]
    header.text = "楚智农购 | 隐私政策"
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_run_font(header.runs[0], size=9, color=MUTED)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer_run = footer.add_run("第 ")
    set_run_font(footer_run, size=9, color=MUTED)
    add_page_number(footer)
    footer_run = footer.add_run(" 页")
    set_run_font(footer_run, size=9, color=MUTED)

    title = document.add_paragraph()
    title.paragraph_format.space_after = Pt(4)
    title_run = title.add_run("楚智农购隐私政策")
    set_run_font(title_run, size=24, color=INK, bold=True)

    subtitle = document.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(2)
    subtitle_run = subtitle.add_run("更新日期：2026年7月25日")
    set_run_font(subtitle_run, size=10.5, color=MUTED)
    subtitle = document.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(14)
    subtitle_run = subtitle.add_run("生效日期：2026年7月25日")
    set_run_font(subtitle_run, size=10.5, color=MUTED)

    add_callout(
        document,
        "特别提示",
        "请在使用楚智农购前仔细阅读本政策。首次启动时，只有在你主动点击“同意并继续”后，"
        "应用才会初始化微信开放平台等第三方 SDK。你可以点击“不同意并退出”，也可以在"
        "“我的 - 设置 - 撤回隐私授权”中撤回同意。",
    )

    document.add_heading("一、适用范围与运营者", level=1)
    add_body(
        document,
        "本政策适用于“楚智农购”Android 移动应用及由该应用提供的商品浏览、手机号登录、"
        "购物车、下单、订单管理、评价、售后、会员、优惠券和消息等服务。",
    )
    add_body(document, "运营者：荆州楚智信息科技有限公司。", bold_prefix="运营者：")
    add_body(
        document,
        "本政策不适用于由第三方独立提供并由其自行控制的服务。对于第三方服务，我们会在"
        "本政策中说明其名称、用途和可能处理的信息。",
    )

    document.add_heading("二、我们如何收集和使用个人信息", level=1)
    document.add_heading("2.1 手机号登录", level=2)
    add_body(
        document,
        "当你获取验证码并登录时，我们会处理手机号和短信验证码，用于验证账号归属、创建"
        "或登录账号及保障账号安全。验证码只用于本次身份验证。",
    )

    document.add_heading("2.2 个人资料", level=2)
    add_body(
        document,
        "你可以自主设置昵称和头像。头像由你通过系统图片选择器主动选取并上传；不提供头像"
        "或自定义昵称，不影响商品浏览等基本功能。",
    )

    document.add_heading("2.3 购物与履约", level=2)
    add_body(
        document,
        "为了生成和履行订单，我们会处理购物车信息、商品与数量、收货人姓名、手机号、收货"
        "地址、订单金额、优惠信息、订单状态、物流信息及你主动填写的订单备注。",
    )
    add_body(
        document,
        "当你发起支付时，我们会将订单号、订单金额和支付所需参数传递给你选择的支付机构，"
        "并接收支付状态。我们不会取得你的银行卡密码或微信支付密码。",
    )

    document.add_heading("2.4 评价与售后", level=2)
    add_body(
        document,
        "当你发布评价或申请售后时，我们会处理评分、文字内容、你主动选择的图片、售后原因、"
        "退款金额及处理进度，用于展示评价、审核申请、处理退换货和退款。每次最多选择三张图片。",
    )

    document.add_heading("2.5 消息与客服", level=2)
    add_body(
        document,
        "我们会根据订单、物流和售后状态向账号内发送消息，并记录已读状态。你主动拨打客服"
        "电话时，应用仅调起系统拨号界面，不读取通话记录。",
    )

    document.add_heading("2.6 网络与安全信息", level=2)
    add_body(
        document,
        "在你访问服务时，服务器可能记录 IP 地址、请求时间、接口地址、响应状态和必要的错误"
        "日志，用于网络连接、故障排查、防止欺诈和保障服务安全。我们不会以实现上述目的为由"
        "读取通讯录、短信内容、精确位置或应用安装列表。",
    )

    document.add_heading("2.7 本地存储", level=2)
    add_body(
        document,
        "应用会在设备本地保存登录令牌、基础账号资料、搜索历史和隐私同意版本，用于保持登录"
        "状态、展示资料、提供搜索记录和避免重复弹窗。这些数据不会由本地存储组件自行上传。",
    )

    add_callout(
        document,
        "当前不收集的信息",
        "楚智农购当前不申请定位、相机、麦克风、通讯录、短信、通话记录或读取全部照片的权限，"
        "也不主动收集 IMEI、IMSI、MAC 地址或已安装应用列表。若后续新增相关功能，我们会先"
        "更新本政策，并在使用具体功能时依法另行征得授权。",
    )

    document.add_heading("三、系统权限与图片选择", level=1)
    add_bullet(document, "网络访问：用于加载商品信息和图片、登录、下单、查询订单及提交售后。")
    add_bullet(
        document,
        "图片选择：头像、评价和售后图片通过 Android 系统图片选择器选择。应用只能读取你本次"
        "主动选择的文件，不申请读取整个相册的广泛权限。",
    )
    add_bullet(
        document,
        "拨号：点击“联系客服”时调起系统拨号界面，由你确认是否拨出；应用不直接拨号，也不读取"
        "通话记录。",
    )

    document.add_heading("四、第三方服务与 SDK", level=1)
    add_body(
        document,
        "我们遵循最小必要原则接入第三方服务。第三方 SDK 只会在你同意本政策后初始化，并在你"
        "使用对应功能时处理实现该功能所需的信息。",
    )

    document.add_heading("4.1 微信开放平台 SDK / 微信支付", level=2)
    add_bullet(document, "提供方：深圳市腾讯计算机系统有限公司、财付通支付科技有限公司。")
    add_bullet(document, "使用目的：检查微信是否可用并调起微信 APP 完成订单支付。")
    add_bullet(
        document,
        "触发场景：同意隐私政策后完成 SDK 注册；只有在你主动选择微信支付时才调起支付能力。",
    )
    add_bullet(
        document,
        "可能处理的信息：网络状态、应用包名和签名、设备与系统基础信息、微信安装状态、订单标识、"
        "订单金额及支付结果。实际处理范围以微信支付规则为准。",
    )
    paragraph = document.add_paragraph()
    run = paragraph.add_run("规则链接：")
    set_run_font(run, bold=True)
    add_hyperlink(
        paragraph,
        "微信支付 OpenSDK 接入说明",
        "https://pay.wechatpay.cn/doc/v3/merchant/4013289321",
    )

    document.add_heading("4.2 阿里云短信服务", level=2)
    add_bullet(document, "提供方：阿里云计算有限公司。")
    add_bullet(document, "使用目的：向你填写的手机号发送登录验证码。")
    add_bullet(document, "处理的信息：手机号、短信模板参数、发送时间和发送结果。")

    document.add_heading("4.3 阿里云对象存储服务", level=2)
    add_bullet(document, "提供方：阿里云计算有限公司。")
    add_bullet(document, "使用目的：存储并分发商品图片、品牌图片和公开法律文件。")
    add_bullet(document, "可能处理的信息：IP 地址、访问时间、请求地址和网络日志。")

    add_body(
        document,
        "Dio、GetX、SharedPreferences、Image Picker、Cached Network Image、WebView 和"
        " DOCX Viewer 等基础开源组件用于应用内部功能实现，不以其自身名义运营账号或向其开发者"
        "发送你的业务数据。",
    )

    document.add_heading("五、共享、转让与公开披露", level=1)
    add_body(
        document,
        "除本政策列明的服务提供商、取得你的单独同意或法律法规另有规定外，我们不会向其他公司、"
        "组织或个人共享个人信息。服务提供商只能按我们的指示为实现约定功能处理必要信息。",
    )
    add_body(
        document,
        "如发生合并、分立、重组、资产转让或类似交易并涉及个人信息转移，我们会要求新的接收方"
        "继续受本政策约束；否则将重新征得你的同意。",
    )
    add_body(
        document,
        "我们原则上不公开披露个人信息。确需公开披露时，将依法说明目的、类型并取得单独同意；"
        "法律法规或有权机关依法要求的除外。",
    )

    document.add_heading("六、存储期限与安全保护", level=1)
    add_body(
        document,
        "我们在中华人民共和国境内存储本应用收集的个人信息，并仅在实现处理目的所必要的最短"
        "期限内保存。账号资料在账号存续期间保存；安全日志按合理期限保存；订单、支付、发票、"
        "退款和售后记录按照电子商务、税务、会计及争议处理等法定要求留存。",
    )
    add_body(
        document,
        "我们采取访问控制、身份认证、传输加密、权限隔离、日志审计和数据备份等合理措施保护"
        "信息安全。互联网服务不存在绝对安全，如发生可能影响你权益的安全事件，我们会依法告知"
        "并采取补救措施。",
    )

    document.add_heading("七、你的权利", level=1)
    add_bullet(document, "访问和更正：在“我的 - 设置”中查看并修改头像和昵称。")
    add_bullet(
        document,
        "删除部分信息：删除收货地址、购物车、收藏和搜索记录；如需删除已发布的评价，可联系客服处理。",
    )
    add_bullet(
        document,
        "撤回同意：在“我的 - 设置 - 撤回隐私授权”中撤回。撤回后应用会退出，下次启动将重新"
        "征求同意；撤回不影响此前基于同意进行处理的合法性。",
    )
    add_bullet(
        document,
        "注销账号：在“我的 - 设置 - 注销账号”中直接注销。账号会立即停用，手机号、头像、昵称、"
        "地址、购物车、收藏、搜索和浏览记录等非交易数据会被删除或匿名化。",
    )
    add_bullet(
        document,
        "无法登录时：访问账号注销申请页面并提交手机号，客服将在核验账号归属后处理。",
    )
    paragraph = document.add_paragraph()
    run = paragraph.add_run("账号注销申请页面：")
    set_run_font(run, bold=True)
    add_hyperlink(
        paragraph,
        "https://vue.climby.cn/api/v2/legal/account-deletion",
        "https://vue.climby.cn/api/v2/legal/account-deletion",
    )
    add_body(
        document,
        "注销后，依法需要留存的订单、支付、退款和售后记录会与正常使用账号隔离，并仅在法定期限"
        "和必要范围内用于履行法定义务、处理争议及保障交易安全。",
    )

    document.add_heading("八、未成年人保护", level=1)
    add_body(
        document,
        "楚智农购主要面向成年人提供农产品购物服务。未满十四周岁的未成年人应在父母或其他"
        "监护人陪同下阅读本政策，并在监护人同意和指导下使用服务。如我们发现未经监护人同意"
        "处理了儿童个人信息，将依法尽快删除或采取其他保护措施。",
    )

    document.add_heading("九、政策更新", level=1)
    add_body(
        document,
        "当业务功能、个人信息处理方式或法律法规发生变化时，我们可能更新本政策。发生重大变化"
        "时，我们会通过弹窗、站内消息或其他显著方式告知；需要重新取得同意的，将在你确认同意后"
        "再按更新后的政策处理信息。",
    )

    document.add_heading("十、联系我们", level=1)
    add_body(document, "运营者：荆州楚智信息科技有限公司。", bold_prefix="运营者：")
    add_body(document, "客服与隐私咨询电话：0716-6666889。", bold_prefix="客服与隐私咨询电话：")
    add_body(
        document,
        "如你对本政策、个人信息处理或账号注销有疑问、意见或投诉，可通过上述电话联系我们。"
        "我们会核验你的身份，并在法律规定的期限内答复。",
    )

    document.add_heading("十一、争议解决", level=1)
    add_body(
        document,
        "本政策适用中华人民共和国法律。因本政策或个人信息处理产生争议的，双方应先友好协商；"
        "协商不成的，可依法向有管辖权的人民法院提起诉讼，或向网信、电信等有关主管部门投诉。",
    )

    document.core_properties.title = "楚智农购隐私政策"
    document.core_properties.subject = "楚智农购 Android 应用个人信息处理规则"
    document.core_properties.author = "荆州楚智信息科技有限公司"
    return document


def main():
    APP_OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    SERVICE_OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document = build_document()
    document.save(APP_OUTPUT)
    copyfile(APP_OUTPUT, SERVICE_OUTPUT)
    print(APP_OUTPUT)
    print(SERVICE_OUTPUT)


if __name__ == "__main__":
    main()
